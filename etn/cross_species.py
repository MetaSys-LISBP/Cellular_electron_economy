"""Cross-species adapters and analysis helpers for electron-transfer networks.

The core ETN chemistry remains model-independent.  This module contains only
model-format adapters, condition configuration helpers, and compact utilities
used by the publication cross-species analyses.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import hashlib
import json
import re
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from cobra import Model, Reaction, Metabolite
from cobra.io import read_sbml_model, load_yaml_model
from cobra.flux_analysis import pfba

from .degree_of_reduction import SAFE_ELEMENTS
from .matching import (
    _metabolite_base_id,
    build_metabolite_table,
    build_matching_context,
    clear_matching_context_cache,
)
from .network_analysis import build_electron_graph, basic_stats, simple_weighted_graph
from .electron_flow_analysis import signed_donor_acceptor_edges
from .flow_decomposition import decompose_source_sink_paths


def sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _split_equation(eq: str):
    for arrow in ("<=>", "=>", "-->", "<--", "<="):
        if arrow in eq:
            left, right = eq.split(arrow, 1)
            return left.strip(), right.strip(), arrow
    raise ValueError(f"No reaction arrow in {eq!r}")


def _parse_side(side: str, key_to_met: dict[str, Metabolite]):
    sto: dict[Metabolite, float] = {}
    if not side.strip():
        return sto
    for term in side.split(" + "):
        term = term.strip()
        match = re.match(r"^([0-9]+(?:\.[0-9]+)?(?:[eE][-+]?\d+)?)\s+(.+)$", term)
        if match:
            coefficient = float(match.group(1))
            key = match.group(2).strip()
        else:
            coefficient, key = 1.0, term
        if key not in key_to_met:
            raise KeyError(f"Unknown metabolite token {key!r}")
        met = key_to_met[key]
        sto[met] = sto.get(met, 0.0) + coefficient
    return sto


def _parse_miriam(value) -> dict[str, list[str]]:
    out: dict[str, list[str]] = defaultdict(list)
    if not value:
        return {}
    for entry in str(value).split(";"):
        if "/" in entry:
            key, val = entry.split("/", 1)
            out[key].append(val)
    return dict(out)


def _sheet_rows(workbook, name: str):
    sheet = workbook[name]
    sheet.reset_dimensions()
    rows = list(sheet.iter_rows(values_only=True))
    return list(rows[0]), rows[1:]


def load_standard_gem_xlsx(path: str | Path, model_id: str) -> Model:
    """Load a RAVEN/standard-GEM Excel workbook used by yeast-GEM/iMT1026.

    The loader intentionally preserves reaction identifiers and bounds while
    constructing a COBRApy model from the METS/RXNS worksheets.
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=True, data_only=True)
    mh, mrows = _sheet_rows(wb, "METS")

    def indices(header, key):
        return [i for i, x in enumerate(header) if str(x).strip().lower() == key.lower()]

    id_i = indices(mh, "ID")[0]
    name_i = indices(mh, "NAME")[0]
    comp_i = indices(mh, "COMPARTMENT")[0]
    repl_i = indices(mh, "REPLACEMENT ID")[0]
    form_i = indices(mh, "COMPOSITION")[0]
    charge_is = indices(mh, "CHARGE")
    miriam_is = indices(mh, "MIRIAM")

    model = Model(model_id)
    keymap: dict[str, Metabolite] = {}
    mets = []
    for row in mrows:
        key, repl = row[id_i], row[repl_i]
        if not key:
            continue
        mid = str(repl or str(key).replace("[", "_").replace("]", ""))
        compartment = str(row[comp_i] or "")
        met = Metabolite(mid, name=str(row[name_i] or key), compartment=compartment)
        met.formula = str(row[form_i]) if row[form_i] else None
        charge = None
        for ci in reversed(charge_is):
            val = row[ci]
            if val is not None and not (isinstance(val, float) and np.isnan(val)):
                charge = val
                break
        if charge is not None:
            try:
                met.charge = int(float(charge))
            except Exception:
                pass
        if miriam_is:
            ann = _parse_miriam(row[miriam_is[0]])
            met.annotation = {k: (v[0] if len(v) == 1 else v) for k, v in ann.items()}
        keymap[str(key)] = met
        mets.append(met)
    model.add_metabolites(mets)

    rh, rrows = _sheet_rows(wb, "RXNS")
    rindex = {str(x).strip().upper(): i for i, x in enumerate(rh)}
    reactions = []
    objective = None
    for row in rrows:
        rid = row[rindex["ID"]]
        equation = row[rindex["EQUATION"]]
        if not rid or not equation:
            continue
        left, right, arrow = _split_equation(str(equation).replace("-->", "=>"))
        rxn = Reaction(str(rid), name=str(row[rindex["NAME"]] or rid))
        default_lb = -1000.0 if arrow == "<=>" else (0.0 if arrow in ("=>", "-->") else -1000.0)
        default_ub = 1000.0 if arrow in ("<=>", "=>", "-->") else 0.0
        lb = row[rindex["LOWER BOUND"]]
        ub = row[rindex["UPPER BOUND"]]
        rxn.bounds = (float(default_lb if lb is None else lb), float(default_ub if ub is None else ub))
        sto = {}
        for met, coef in _parse_side(left, keymap).items():
            sto[met] = sto.get(met, 0.0) - coef
        for met, coef in _parse_side(right, keymap).items():
            sto[met] = sto.get(met, 0.0) + coef
        rxn.add_metabolites(sto)
        if "SUBSYSTEM" in rindex:
            rxn.subsystem = str(row[rindex["SUBSYSTEM"]] or "")
        if "GENE ASSOCIATION" in rindex:
            gpr = str(row[rindex["GENE ASSOCIATION"]] or "")
            if gpr:
                try:
                    rxn.gene_reaction_rule = gpr.replace(";", " or ")
                except Exception:
                    pass
        if "MIRIAM" in rindex:
            ann = _parse_miriam(row[rindex["MIRIAM"]])
            rxn.annotation = {k: (v[0] if len(v) == 1 else v) for k, v in ann.items()}
        reactions.append(rxn)
        if "OBJECTIVE" in rindex and row[rindex["OBJECTIVE"]] not in (None, 0, 0.0):
            objective = str(rid)
    model.add_reactions(reactions)
    if objective:
        model.objective = objective
    return model


def load_bacillus_xlsx(path: str | Path) -> Model:
    """Load iBB1018 from its publication Excel workbook."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Metabolite List"]
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    idx = {x: i for i, x in enumerate(h)}
    model = Model("iBB1018")
    mets = {}
    for row in rows[1:]:
        mid = row[idx["Abbreviation"]]
        if not mid:
            continue
        raw_comp = str(row[idx["Compartment"]] or "c").lower()
        comp = "c" if raw_comp == "cytosol" else "e" if raw_comp in {"extracellular", "extracellular space"} else raw_comp
        met = Metabolite(str(mid), name=str(row[idx["Description"]] or mid), compartment=comp)
        met.formula = str(row[idx["Formula"]]) if row[idx["Formula"]] else None
        met.charge = int(row[idx["Charge"]]) if row[idx["Charge"]] is not None else None
        mets[mid] = met
    model.add_metabolites(list(mets.values()))

    ws = wb["Reaction List"]
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    idx = {x: i for i, x in enumerate(h)}
    reactions = []
    objective = None
    for row in rows[1:]:
        rid = row[idx["Abbreviation"]]
        if not rid:
            continue
        left, right, _ = _split_equation(str(row[idx["Reaction"]]).replace("-->", "=>"))
        rxn = Reaction(str(rid), name=str(row[idx["Description"]] or rid))
        rxn.bounds = (float(row[idx["Lower bound"]]), float(row[idx["Upper bound"]]))
        sto = {}
        for met, coef in _parse_side(left, mets).items():
            sto[met] = sto.get(met, 0.0) - coef
        for met, coef in _parse_side(right, mets).items():
            sto[met] = sto.get(met, 0.0) + coef
        rxn.add_metabolites(sto)
        rxn.subsystem = str(row[idx["Subsystem"]] or "")
        gpr = str(row[idx["GPR"]] or "")
        if gpr:
            rxn.gene_reaction_rule = gpr
        reactions.append(rxn)
        if row[idx["Objective"]] not in (None, 0, 0.0):
            objective = str(rid)
    model.add_reactions(reactions)
    if objective:
        model.objective = objective
    return model


def normalize_yeast_for_etn(model: Model) -> Model:
    """Return an isostoichiometric copy with BiGG-like metabolite IDs where annotated."""
    new = Model(model.id + "_etn")
    old_to_new = {}
    used = set()
    for met in model.metabolites:
        bigg = met.annotation.get("bigg.metabolite") if hasattr(met, "annotation") else None
        if isinstance(bigg, list):
            bigg = bigg[0] if bigg else None
        nid = f"{bigg}_{met.compartment}" if bigg else f"{met.id}_{met.compartment}"
        if nid in used:
            nid = f"{met.id}_{met.compartment}"
        used.add(nid)
        new_met = Metabolite(nid, name=met.name, compartment=met.compartment)
        new_met.formula, new_met.charge = met.formula, met.charge
        new_met.annotation = dict(met.annotation)
        old_to_new[met.id] = new_met
    new.add_metabolites(list(old_to_new.values()))
    for rxn in model.reactions:
        new_rxn = Reaction(rxn.id, name=rxn.name)
        new_rxn.bounds = rxn.bounds
        new_rxn.subsystem = rxn.subsystem
        new_rxn.annotation = dict(rxn.annotation)
        new_rxn.add_metabolites({old_to_new[m.id]: c for m, c in rxn.metabolites.items()})
        new.add_reactions([new_rxn])
    return new


def chemical_model(model: Model, *, yeast: bool = False) -> Model:
    """Remove structural pseudo-reactions before chemistry classification."""
    out = model.copy()
    remove = []
    for rxn in out.reactions:
        rid = rxn.id.lower()
        subsystem = (rxn.subsystem or "").lower()
        name = (rxn.name or "").lower()
        sbo = str(rxn.annotation.get("sbo", "")) if hasattr(rxn, "annotation") else ""
        structural = (
            rid.startswith(("ex_", "dm_", "sk_"))
            or (rid.startswith("ex") and "exchange" in name)
            or "exchange reaction" in subsystem
            or rid.startswith("biomass")
            or rid == "growth"
            or name.strip() in {"biomass composition (g/g)", "growth"}
            or "biomass composition" in subsystem
            or "pseudoreaction" in name
            or (yeast and ("sbo:0000627" in sbo or "sbo:0000632" in sbo or subsystem == "growth"))
        )
        if structural:
            remove.append(rxn)
    out.remove_reactions(remove, remove_orphans=False)
    return out


def _unscorable_reason(met: Metabolite) -> str:
    if not met.formula:
        return "missing_formula"
    try:
        elements = met.elements or {}
    except Exception:
        return "invalid_formula"
    bad = set(elements) - SAFE_ELEMENTS
    if bad:
        if bad & {"Fe", "Cu", "Mo", "Co", "Ni", "Mn", "Zn", "Ag", "Cd", "Hg", "W"}:
            return "transition_metal"
        if bad & {"R", "X"}:
            return "generic_placeholder"
        return "unsupported_element"
    return "other"


def classify_and_qc(full_model: Model, chem_model: Model, organism: str):
    # Matching contexts are cached by model identity.  A cross-species batch clears
    # the cache once before reconstruction, then retains one context per chemical
    # model so condition-specific electron-flux calculation reuses the exact reconstructed ETN.
    mt = build_metabolite_table(chem_model)
    context = build_matching_context(chem_model, mt)
    classification = [context.results_by_reaction[r.id] for r in chem_model.reactions]
    count = Counter(row["redox_status"] for row in classification)
    redox = [row for row in classification if row.get("is_redox")]
    resolution = Counter(row.get("resolution_status") for row in redox)
    graph = build_electron_graph(chem_model, classification)
    stats = basic_stats(graph) if graph.number_of_nodes() else {
        "n_nodes": 0, "n_edges_simple": 0, "n_edges_multi": 0, "n_reactions": 0,
        "n_weakly_connected_components": 0, "largest_component_size": 0,
    }
    pairs = set()
    for row in redox:
        for pair in row.get("pairs", []):
            pairs.add(tuple(sorted((_metabolite_base_id(pair["reactant"]), _metabolite_base_id(pair["product"])))))
    reasons = Counter()
    for row in classification:
        if row.get("redox_status") == "unscorable":
            rxn = chem_model.reactions.get_by_id(row["reaction_id"])
            for met in rxn.metabolites:
                if not mt[met.id].ok:
                    reasons[_unscorable_reason(met)] += 1
    analyzable = len(chem_model.reactions)
    classified = count.get("redox", 0) + count.get("non_redox", 0)
    qc = {
        "organism": organism,
        "total_reactions": len(full_model.reactions),
        "total_metabolites": len(full_model.metabolites),
        "total_genes": len(full_model.genes),
        "chemical_reactions": analyzable,
        "confident_redox": count.get("redox", 0),
        "confident_nonredox": count.get("non_redox", 0),
        "ambiguous": count.get("ambiguous", 0),
        "unscorable": count.get("unscorable", 0),
        "classified_fraction": classified / analyzable if analyzable else np.nan,
        "fully_resolved_redox": resolution.get("fully_resolved", 0),
        "redox_total": len(redox),
        "fully_resolved_fraction": resolution.get("fully_resolved", 0) / len(redox) if redox else np.nan,
        "distinct_redox_pairs": len(pairs),
        "etn_nodes": stats["n_nodes"],
        "etn_edges": stats["n_edges_simple"],
        "etn_reaction_edges": stats["n_edges_multi"],
        "etn_reactions": stats["n_reactions"],
        "wcc_count": stats["n_weakly_connected_components"],
        "largest_wcc": stats["largest_component_size"],
        "largest_wcc_fraction": stats["largest_component_size"] / stats["n_nodes"] if stats["n_nodes"] else np.nan,
        "unscorable_reasons": json.dumps(dict(reasons), sort_keys=True),
    }
    simple = simple_weighted_graph(graph)
    hubs = []
    for node in simple.nodes:
        degree = sum(d.get("weight", 1) for *_, d in simple.in_edges(node, data=True)) + sum(
            d.get("weight", 1) for *_, d in simple.out_edges(node, data=True)
        )
        name = chem_model.metabolites.get_by_id(node).name if node in chem_model.metabolites else node
        hubs.append((node, name, degree))
    hubs.sort(key=lambda x: x[2], reverse=True)
    return qc, classification, graph, hubs[:15]


def etn_metrics(chem_model: Model, solution, substrate_uptake: float):
    mt = build_metabolite_table(chem_model)
    # Matching contexts are cached by model identity. Reusing them across
    # physiological states avoids repeating the expensive chemical reconstruction.
    build_matching_context(chem_model, mt)
    edges = signed_donor_acceptor_edges(chem_model, mt, solution)
    _, summary, stats = decompose_source_sink_paths(edges)
    return {
        "cumulative_e_flux": float(summary["total_edge_activity"]),
        "net_e_flux": float(summary["net_source_flux"]),
        "effective_transfer_depth": float(summary["effective_transfer_depth"]),
        "cumulative_per_glucose": float(summary["total_edge_activity"] / substrate_uptake),
        "net_per_glucose": float(summary["net_source_flux"] / substrate_uptake),
    }, edges, stats


def relay_by_base(stats_df: pd.DataFrame, carriers=None):
    carriers = carriers or ["nadh", "nadph", "fadh2", "q6h2", "q8h2", "mql7", "mql8", "2dmmql8"]
    rows = []
    for base in carriers:
        selected = stats_df[stats_df["node"].map(lambda x: _metabolite_base_id(str(x)) == base)]
        rows.append({"carrier": base, "relay": float(selected["relay_throughflow"].sum()) if len(selected) else 0.0})
    return rows


def apply_yeast_anaerobic(model: Model, amino_acid_tsv: str | Path) -> Model:
    """Apply the yeast-GEM v9.1.0 model-provided anaerobic edits."""
    def setcoef(rxn, met, value):
        current = rxn.metabolites.get(met, 0.0)
        delta = float(value) - current
        if delta:
            rxn.add_metabolites({met: delta}, combine=True)

    def rebalance(rxn):
        proton = model.metabolites.get_by_id("s_0794")
        setcoef(rxn, proton, 0.0)
        imbalance = sum((m.charge or 0.0) * c for m, c in rxn.metabolites.items())
        setcoef(rxn, proton, -imbalance)

    cof = model.reactions.get_by_id("r_4598")
    setcoef(cof, model.metabolites.get_by_id("s_3714"), 0.0)
    rebalance(cof)

    protein = model.reactions.get_by_id("r_4047")
    weights = {"C": 12.01, "H": 1.008, "N": 14.007, "O": 15.999, "P": 30.974, "S": 32.06,
               "R": 0.0, "Fe": 55.845, "K": 39.098, "Na": 22.99, "Cl": 35.45, "Mn": 54.938,
               "Zn": 65.38, "Ca": 40.078, "Mg": 24.305, "Cu": 63.546}

    def fw(met):
        total = 0.0
        for element, number in re.findall(r"([A-Z][a-z]*)(\d*)", met.formula or ""):
            total += (int(number) if number else 1) * weights[element]
        return total

    def protein_mass():
        return sum(-coef * (fw(met) - 2.016) for met, coef in protein.metabolites.items() if coef < 0) / 1000.0

    target = protein_mass()
    aa = pd.read_csv(amino_acid_tsv, sep="\t", header=None, skiprows=1)
    for _, row in aa.iterrows():
        ratio = float(row.iloc[5])
        sm = model.metabolites.get_by_id(row.iloc[1])
        pm = model.metabolites.get_by_id(row.iloc[2])
        setcoef(protein, sm, -ratio)
        setcoef(protein, pm, ratio)
    factor = target / protein_mass()
    changes = {met: (factor - 1) * coef for met, coef in list(protein.metabolites.items()) if met.name != "protein"}
    if changes:
        protein.add_metabolites(changes, combine=True)
    rebalance(protein)

    for rid, lower_bound in {
        "r_1992": 0.0, "r_1757": -1000.0, "r_1915": -1000.0, "r_2106": -1000.0,
        "r_2134": -1000.0, "r_1994": -1000.0, "r_2189": -1000.0, "r_2137": 0.0,
        "r_1967": -1000.0, "r_1548": -1000.0,
    }.items():
        model.reactions.get_by_id(rid).lower_bound = lower_bound
    for rid in ("r_0714", "r_0659"):
        model.reactions.get_by_id(rid).bounds = (0, 0)
    biomass = model.reactions.get_by_id("r_4041")
    biomass.add_metabolites({
        model.metabolites.get_by_id("s_0689"): 0.08,
        model.metabolites.get_by_id("s_0687"): -0.08,
        model.metabolites.get_by_id("s_0794"): -0.16,
    }, combine=True)
    return model


def solve_growth_pfba(model: Model, objective: str):
    model.objective = objective
    optimum = model.optimize()
    if optimum.status != "optimal" or optimum.objective_value is None or optimum.objective_value <= 1e-10:
        raise RuntimeError(f"No positive growth optimum for {model.id}")
    solution = pfba(model, fraction_of_optimum=1.0)
    if solution.status != "optimal":
        raise RuntimeError(f"pFBA failed for {model.id}")
    return solution


def load_cross_species_models(data_dir: str | Path):
    data_dir = Path(data_dir)
    yeast = load_standard_gem_xlsx(data_dir / "yeast-GEM-v9.1.0.xlsx", "yeastGEM_v9.1.0")
    bacillus = load_bacillus_xlsx(data_dir / "iBB1018.xlsx")
    salmonella = read_sbml_model(str(data_dir / "iYS1720.xml.gz"))
    pichia = load_standard_gem_xlsx(data_dir / "iMT1026_model.xlsx", "iMT1026")
    ijo1366 = read_sbml_model(str(data_dir / "iJO1366.xml.gz"))
    return {
        "S. cerevisiae": yeast,
        "B. subtilis": bacillus,
        "S. enterica": salmonella,
        "K. phaffii": pichia,
        "iJO1366": ijo1366,
    }
