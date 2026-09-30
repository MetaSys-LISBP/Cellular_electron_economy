"""
make_excel_table.py
--------------------
Builds a professionally formatted Excel workbook listing every
electron-transfer (redox) reaction identified in the reconstructed
network, for readers who prefer a spreadsheet to the TSV in results/.

Sheet 1 "Redox reactions": all reactions with redox_status == "redox"
  -- id, name, full equation, reversibility, electron donor(s)/acceptor(s)
  (both BiGG IDs and readable names), number of electrons transferred,
  and the reaction's metabolic subsystem.
Sheet 2 "All reactions (reference)": the complete, unfiltered table (all
  2,712 reactions of iML1515) with their redox_status, for provenance/
  reproducibility -- so every number in Sheet 1 can be traced back to the
  full reaction set it was filtered from.
Sheet 3 "Legend": column definitions and summary counts.

Usage: python scripts/make_excel_table.py (after run_analysis.py)
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RESULTS = os.path.join(REPO_ROOT, "results")

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill(start_color="2A5C8A", end_color="2A5C8A", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
REDOX_FILL = PatternFill(start_color="E8F1E4", end_color="E8F1E4", fill_type="solid")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMN_RENAME = {
    "reaction_id": "Reaction ID",
    "reaction_name": "Reaction name",
    "reaction_equation": "Full reaction (readable)",
    "reversible": "Reversible",
    "redox_status": "Redox status",
    "n_electrons_transferred": "Electrons transferred",
    "donor_metabolite_ids": "Electron donor (BiGG ID)",
    "donor_metabolite_names": "Electron donor (name)",
    "acceptor_metabolite_ids": "Electron acceptor (BiGG ID)",
    "acceptor_metabolite_names": "Electron acceptor (name)",
    "subsystem": "Subsystem",
}
COLUMN_WIDTHS = {
    "Reaction ID": 16, "Reaction name": 34, "Full reaction (readable)": 70,
    "Reversible": 11, "Redox status": 12, "Electrons transferred": 12,
    "Electron donor (BiGG ID)": 16, "Electron donor (name)": 26,
    "Electron acceptor (BiGG ID)": 16, "Electron acceptor (name)": 26,
    "Subsystem": 30,
}


def _write_sheet(ws, df: pd.DataFrame, freeze_after_header=True):
    headers = list(df.columns)
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = COLUMN_WIDTHS.get(h, 18)
    ws.row_dimensions[1].height = 28

    is_redox_col = headers.index("Redox status") + 1 if "Redox status" in headers else None
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=("" if pd.isna(val) else val))
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=(headers[j - 1] == "Full reaction (readable)"))
        if is_redox_col and row[is_redox_col - 1] == "redox":
            for j in range(1, len(headers) + 1):
                ws.cell(row=i, column=j).fill = REDOX_FILL

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"
    if freeze_after_header:
        ws.freeze_panes = "A2"


def build_workbook(output_path: str):
    et = pd.read_csv(os.path.join(RESULTS, "electron_transfer_table.tsv"), sep="\t")
    et_renamed = et.rename(columns=COLUMN_RENAME)

    redox_only = et_renamed[et_renamed["Redox status"] == "redox"].copy()
    redox_only = redox_only.sort_values("Electrons transferred", ascending=False).reset_index(drop=True)

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Redox reactions"
    _write_sheet(ws1, redox_only)

    ws2 = wb.create_sheet("All reactions (reference)")
    _write_sheet(ws2, et_renamed)

    ws3 = wb.create_sheet("Legend")
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 90
    rows = [
        ("Column", "Definition"),
        ("Reaction ID", "Reaction identifier in the iML1515 genome-scale model (BiGG namespace)."),
        ("Reaction name", "Human-readable reaction name from the model."),
        ("Full reaction (readable)", "Complete stoichiometric equation with metabolite names (not IDs)."),
        ("Reversible", "Whether the reaction's flux bounds in iML1515 allow both directions."),
        ("Redox status", "'redox' = an electron-transfer reaction was identified; 'non_redox' = classified, no transfer; "
                          "'missing_formula' = a participating metabolite contains an element outside the supported set "
                          "(transition metal, redox-active metalloid, or generic R/X group); 'exchange_reaction' = "
                          "transport/exchange reaction (excluded a priori); 'pseudo_reaction' = biomass objective function."),
        ("Electrons transferred", "Number of electrons moved from the donor to the acceptor in one reaction turnover, "
                                   "computed from the change in reduction degree (gamma) of the matched donor/acceptor pair "
                                   "(Roels, 1980)."),
        ("Electron donor (BiGG ID / name)", "The metabolite that is oxidised (loses electrons) in this reaction."),
        ("Electron acceptor (BiGG ID / name)", "The metabolite that is reduced (gains electrons) in this reaction."),
        ("Subsystem", "Curated metabolic subsystem/pathway annotation from the iML1515 model."),
        ("", ""),
        ("Summary", ""),
        ("Total reactions in iML1515", str(len(et))),
        ("Redox (electron-transfer) reactions identified", str(len(redox_only))),
        ("Non-redox reactions", str((et.redox_status == "non_redox").sum())),
        ("Excluded: exchange/transport reactions", str((et.redox_status == "exchange_reaction").sum())),
        ("Excluded: missing formula (metal/metalloid/generic group)", str((et.redox_status == "missing_formula").sum())),
        ("Excluded: biomass pseudo-reaction", str((et.redox_status == "pseudo_reaction").sum())),
        ("", ""),
        ("Source", "Generated by scripts/run_analysis.py + scripts/make_excel_table.py; see README.md to reproduce."),
        ("Model", "iML1515 (Monk et al. 2017, Nat Biotechnol 35:904-908), bundled at data/iML1515.xml."),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        ca = ws3.cell(row=i, column=1, value=a)
        cb = ws3.cell(row=i, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1 or a == "Summary":
            ca.font = Font(name="Arial", bold=True, size=11)
            cb.font = Font(name="Arial", bold=True, size=11)
        else:
            ca.font = Font(name="Arial", bold=(b == "" and a != ""), size=10)
            cb.font = Font(name="Arial", size=10)

    wb.save(output_path)
    print(f"[make_excel_table] Wrote {output_path} "
          f"({len(redox_only)} redox reactions, {len(et_renamed)} total reactions)")


def build_pairs_workbook(output_path: str):
    """Second workbook: every confidently accepted matched electron-transfer
    pair network-wide, one row per reaction x pair."""
    match = pd.read_csv(os.path.join(RESULTS, "matched_pairs_all.csv"))
    met = pd.read_csv(os.path.join(RESULTS, "metabolites_ecoli.csv"))
    name_map = dict(zip(met["met_id"], met["name"]))
    match = match.rename(columns={
        "reaction_id": "Reaction ID", "reactant": "Pair member A (BiGG ID)", "product": "Pair member B (BiGG ID)",
        "source": "Match origin", "electrons": "Electrons transferred",
    })
    match["Pair member A (name)"] = match["Pair member A (BiGG ID)"].map(name_map)
    match["Pair member B (name)"] = match["Pair member B (BiGG ID)"].map(name_map)
    match["Match origin"] = match["Match origin"].map({
        "curated_cofactor_list": "Curated seed list", "ab_initio_formula_match": "Ab initio (formula matching)"})
    cols = ["Reaction ID", "Pair member A (BiGG ID)", "Pair member A (name)", "Pair member B (BiGG ID)",
            "Pair member B (name)", "Electrons transferred", "Match origin"]
    match = match[cols].sort_values(["Match origin", "Reaction ID"]).reset_index(drop=True)

    widths = {"Reaction ID": 16, "Pair member A (BiGG ID)": 16, "Pair member A (name)": 30,
              "Pair member B (BiGG ID)": 16, "Pair member B (name)": 30,
              "Electrons transferred": 14, "Match origin": 24}
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "All matched pairs"
    for j, h in enumerate(cols, start=1):
        c = ws1.cell(row=1, column=j, value=h)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws1.column_dimensions[get_column_letter(j)].width = widths.get(h, 18)
    ws1.row_dimensions[1].height = 26
    for i, row in enumerate(match.itertuples(index=False), start=2):
        for j, val in enumerate(row, start=1):
            c = ws1.cell(row=i, column=j, value=("" if pd.isna(val) else val))
            c.font = BODY_FONT; c.border = BORDER
        if match.iloc[i - 2]["Match origin"] == "Ab initio (formula matching)":
            for j in range(1, len(cols) + 1):
                ws1.cell(row=i, column=j).fill = REDOX_FILL
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(match) + 1}"
    ws1.freeze_panes = "A2"

    # Sheet 2: distinct ab initio couples (deduplicated across reactions)
    novel = pd.read_csv(os.path.join(RESULTS, "ab_initio_redox_pairs.csv"))
    novel = novel.rename(columns={
        "reactant": "Pair member A (BiGG ID)", "product": "Pair member B (BiGG ID)",
        "reactant_name": "Pair member A (name)", "product_name": "Pair member B (name)",
        "n_reactions": "# reactions using this pair", "reactions": "Reaction IDs",
        "mean_electrons": "Mean electrons transferred",
    })
    ws2 = wb.create_sheet("Distinct ab initio couples")
    cols2 = ["Pair member A (BiGG ID)", "Pair member A (name)", "Pair member B (BiGG ID)", "Pair member B (name)",
             "# reactions using this pair", "Mean electrons transferred", "Reaction IDs"]
    widths2 = {"Pair member A (BiGG ID)": 16, "Pair member A (name)": 30, "Pair member B (BiGG ID)": 16,
               "Pair member B (name)": 30, "# reactions using this pair": 14,
               "Mean electrons transferred": 14, "Reaction IDs": 40}
    for j, h in enumerate(cols2, start=1):
        c = ws2.cell(row=1, column=j, value=h)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws2.column_dimensions[get_column_letter(j)].width = widths2.get(h, 18)
    ws2.row_dimensions[1].height = 26
    novel_sorted = novel.sort_values("# reactions using this pair", ascending=False)[cols2].reset_index(drop=True)
    for i, row in enumerate(novel_sorted.itertuples(index=False), start=2):
        for j, val in enumerate(row, start=1):
            c = ws2.cell(row=i, column=j, value=("" if pd.isna(val) else val))
            c.font = BODY_FONT; c.border = BORDER
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(cols2))}{len(novel_sorted) + 1}"
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Legend")
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 90
    rows = [
        ("Column", "Definition"),
        ("Reaction ID", "Reaction identifier in iML1515 (BiGG namespace) in which this pair was matched."),
        ("Pair members A/B (BiGG ID / name)", "The reactant-side and product-side members of the matched redox couple in the reaction as written; donor/acceptor role is determined from Δγ and, in flux analyses, from flux direction."),
        ("Electrons transferred", "|delta-gamma| for this specific pair (Roels reduction degree)."),
        ("Match origin", "'Curated seed list' = one of 16 universal cofactor pairs supplied to the algorithm; "
                          "'Ab initio (formula matching)' = inferred without direct seeding from identical C/N/S/P atom "
                          "counts, no redox information supplied (shaded rows on sheet 1)."),
        ("", ""),
        ("Sheet 2", f"The {len(novel)} distinct ab initio metabolite-level couples (Sheet 1 deduplicated across reactions), with usage counts."),
        ("", ""),
        ("Summary", ""),
        ("Total matched pairs (Sheet 1)", str(len(match))),
        ("  - from curated seed list", str((match["Match origin"] == "Curated seed list").sum())),
        ("  - inferred ab initio", str((match["Match origin"] == "Ab initio (formula matching)").sum())),
        ("Distinct ab initio couples (Sheet 2)", str(len(novel))),
        ("Model", "iML1515 (Monk et al. 2017, Nat Biotechnol 35:904-908), bundled at data/iML1515.xml."),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        ca = ws3.cell(row=i, column=1, value=a)
        cb = ws3.cell(row=i, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.font = Font(name="Arial", bold=(i == 1 or a == "Summary"), size=10)
        cb.font = Font(name="Arial", bold=(i == 1 or a == "Summary"), size=10)

    wb.save(output_path)
    print(f"[make_excel_table] Wrote {output_path} "
          f"({len(match)} matched pairs, {len(novel)} distinct ab initio couples)")


if __name__ == "__main__":
    build_workbook(os.path.join(RESULTS, "redox_reactions_iML1515.xlsx"))
    build_pairs_workbook(os.path.join(RESULTS, "redox_pairs_iML1515.xlsx"))
